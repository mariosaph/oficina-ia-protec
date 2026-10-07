#!/usr/bin/env python3
"""
preencher_oficio.py - preenche o modelo de ofício do SEI da UFAM (.docx) a partir de um arquivo .json.

COMO RODAR (no Windows troque python3 por py)
  python3 robos/preencher_oficio.py --exemplo                 escreve saida/oficio-exemplo.json para editar
  python3 robos/preencher_oficio.py saida/oficio.json         gera saida/oficio-<assunto>.docx

O QUE ELE FAZ
  Troca cada [CAMPO] do modelo pelo valor do .json, sem mexer na formatação. O que não estiver no .json
  continua no arquivo como [CAMPO], marcado em amarelo, e aparece na lista "Para conferir". Valor que
  traz [CONFERIR: ...] também fica em amarelo. O bloco de assinatura eletrônica, o código verificador e o
  CRC ficam como estão: o SEI preenche quando você assina.

  Modelo: .claude/skills/redigir-oficio/modelo-oficio-sei-ufam.docx (ou --modelo outro.docx)
  Precisa de: pip install python-docx

CÓDIGOS DE SAÍDA
  0 gerou o arquivo | 1 erro (arquivo, .json ou biblioteca faltando)
"""
import argparse
import copy
import json
import re
import sys
import unicodedata
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
MODELO = RAIZ / ".claude" / "skills" / "redigir-oficio" / "modelo-oficio-sei-ufam.docx"
SAIDA = RAIZ / "saida"
if not MODELO.exists() and (AQUI / MODELO.name).exists():
    # zip da skill no claude.ai ou plugin da loja da oficina: o modelo fica ao lado do script,
    # e a saída vai para a pasta onde a pessoa está (não para dentro da skill)
    MODELO = AQUI / MODELO.name
    SAIDA = Path.cwd() / "saida"

# chave do .json -> marcador no modelo (o gerador do modelo, _build/oficio_sei.py, lê esta tabela)
CAMPOS = {
    "unidade": "[NOME DA UNIDADE]",
    "numero": "[NÚMERO]",
    "ano": "[ANO]",
    "sigla": "[SIGLA DA UNIDADE]",
    "cidade": "[CIDADE]",
    "data": "[DIA DE MÊS DE ANO]",
    "tratamento": "[TRATAMENTO: Ao Senhor, À Senhora, Ao Diretor...]",
    "destinatario": "[NOME OU CARGO DO DESTINATÁRIO]",
    "orgao": "[ÓRGÃO OU EMPRESA]",
    "endereco": "[LOGRADOURO, NÚMERO, BAIRRO]",
    "cep_cidade": "[CEP 00000-000, CIDADE/UF]",
    "assunto": "[ASSUNTO EM POUCAS PALAVRAS]",
    "vocativo": "[VOCATIVO: Senhor Diretor]",
    "fecho": "[Atenciosamente ou Respeitosamente]",
    "signatario": "[NOME DO SIGNATÁRIO]",
    "cargo": "[CARGO DO SIGNATÁRIO]",
    "endereco_unidade": "[ENDEREÇO DA UNIDADE]",
    "telefone_unidade": "[TELEFONE DA UNIDADE]",
    "cep_cidade_unidade": "[CEP E CIDADE DA UNIDADE]",
    "email_unidade": "[E-MAIL DA UNIDADE]",
    "processo": "[23105.000000/2026-00]",
    "sei": "[0000000]",
}
# parágrafos que se repetem: um por item da lista do .json
LISTAS = {
    "paragrafos": "[PARÁGRAFO DO TEXTO]",
    "itens": "[ITEM NUMERADO]",
    "paragrafos_finais": "[PARÁGRAFO DE ENCERRAMENTO]",
}
OPCIONAIS = {"itens", "paragrafos_finais"}   # fora do .json: o trecho sai do ofício
NOTA = "[NOTA DO MODELO"                       # parágrafo de instrução: sai do arquivo gerado

EXEMPLO = {
    "unidade": "Pró-Reitoria de Inovação Tecnológica",
    "numero": "[CONFERIR: o SEI gera o número]",
    "ano": "2026",
    "sigla": "PROTEC",
    "cidade": "Manaus",
    "data": "6 de outubro de 2026",
    "tratamento": "Ao Senhor",
    "destinatario": "[CONFERIR: nome do diretor]",
    "orgao": "Diretor da Empresa X",
    "endereco": "[CONFERIR: endereço da empresa]",
    "cep_cidade": "[CONFERIR: CEP e cidade/UF]",
    "assunto": "Convite para reunião de cooperação técnica",
    "vocativo": "Senhor Diretor",
    "paragrafos": [
        "Convido Vossa Senhoria para uma reunião sobre cooperação técnica entre a Universidade Federal do Amazonas e a Empresa X, na semana de 19 de outubro de 2026, em data e horário a combinar."
    ],
    "paragrafos_finais": [
        "Coloco-me à disposição para alinhar a pauta pelo e-mail [CONFERIR: e-mail da unidade]."
    ],
    "fecho": "Atenciosamente",
    "signatario": "[CONFERIR: nome de quem assina]",
    "cargo": "[CONFERIR: cargo de quem assina]",
}


def falhar(msg):
    print(msg)
    sys.exit(1)


def pedacos(doc):
    """Todos os parágrafos: corpo, tabelas e rodapés/cabeçalhos."""
    def de(bloco):
        for p in bloco.paragraphs:
            yield p
        for t in bloco.tables:
            for linha in t.rows:
                for cel in linha.cells:
                    yield from de(cel)
    yield from de(doc)
    for s in doc.sections:
        for parte in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            if parte is not None:
                yield from de(parte)


def marcar(run, amarelo):
    from docx.enum.text import WD_COLOR_INDEX
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW if amarelo else None


def definir(run, valor):
    """Põe o valor no trecho. Só os [CONFERIR: ...] ficam em amarelo: o trecho é dividido em pedaços
    com a mesma formatação."""
    pedacos_txt = [p for p in re.split(r"(\[CONFERIR[^\]]*\])", valor) if p]
    if not pedacos_txt:
        run.text = ""
        marcar(run, False)
        return
    run.text = pedacos_txt[0]
    marcar(run, pedacos_txt[0].startswith("[CONFERIR"))
    ancora = run._r
    for p in pedacos_txt[1:]:
        novo = copy.deepcopy(run._r)
        ancora.addnext(novo)
        ancora = novo
        from docx.text.run import Run
        r = Run(novo, run._parent)
        r.text = p
        marcar(r, p.startswith("[CONFERIR"))


def apagar(par):
    el = par._element
    el.getparent().remove(el)


def slug(txt):
    t = unicodedata.normalize("NFD", txt or "oficio").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:50] or "oficio"


def preencher(dados, modelo):
    try:
        from docx import Document
    except ImportError:
        falhar("Falta a biblioteca: pip install python-docx")
    if not modelo.exists():
        falhar(f"Não achei o modelo {modelo}.")
    doc = Document(str(modelo))
    marcador_para_chave = {v: k for k, v in CAMPOS.items()}

    # 1. listas: clona o parágrafo-modelo para cada item
    for chave, marcador in LISTAS.items():
        alvos = [p for p in pedacos(doc) if any(r.text == marcador for r in p.runs)]
        for par in alvos:
            valores = dados.get(chave)
            if valores is None:
                if chave in OPCIONAIS:
                    apagar(par)
                continue
            if isinstance(valores, str):
                valores = [valores]
            if not valores:
                apagar(par)
                continue
            ancora = par._element
            for i, valor in enumerate(valores, 1):
                novo = copy.deepcopy(par._element)
                ancora.addnext(novo)
                ancora = novo
                from docx.text.paragraph import Paragraph
                p2 = Paragraph(novo, par._parent)
                for r in list(p2.runs):
                    if r.text == marcador:
                        definir(r, str(valor))
                    elif chave == "itens" and re.fullmatch(r"\d+\.\s", r.text):
                        r.text = f"{i}. "
            apagar(par)

    # 2. campos de uma linha
    for par in list(pedacos(doc)):
        if any(r.text.startswith(NOTA) for r in par.runs):
            apagar(par)
            continue
        for r in list(par.runs):
            chave = marcador_para_chave.get(r.text)
            if chave is None or chave not in dados:
                continue
            valor = str(dados[chave])
            if valor == "" and all(x.text in (r.text, "") for x in par.runs):
                apagar(par)          # linha opcional deixada em branco (ex.: endereço)
                break
            definir(r, valor)

    # 3. o que ficou para conferir
    conferir = []
    for par in pedacos(doc):
        for r in par.runs:
            if r.text in marcador_para_chave or r.text in LISTAS.values():
                conferir.append(r.text)
            elif "[CONFERIR" in r.text:
                conferir += re.findall(r"\[CONFERIR[^\]]*\]", r.text)
    return doc, list(dict.fromkeys(conferir))


def main():
    ap = argparse.ArgumentParser(description="Preenche o modelo de ofício do SEI da UFAM.")
    ap.add_argument("dados", nargs="?", help="arquivo .json com os campos do ofício")
    ap.add_argument("--modelo", type=Path, default=MODELO)
    ap.add_argument("--saida", type=Path, help="nome do .docx gerado (padrão: saida/oficio-<assunto>.docx)")
    ap.add_argument("--exemplo", action="store_true", help="escreve saida/oficio-exemplo.json e sai")
    op = ap.parse_args()

    SAIDA.mkdir(exist_ok=True)
    if op.exemplo:
        destino = SAIDA / "oficio-exemplo.json"
        destino.write_text(json.dumps(EXEMPLO, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Exemplo em {destino}. Campos aceitos: {', '.join([*CAMPOS, *LISTAS])}")
        return
    if not op.dados:
        falhar("Diga o arquivo .json (ou rode com --exemplo para ver um).")
    try:
        dados = json.loads(Path(op.dados).read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as erro:
        falhar(f"Não consegui ler {op.dados}: {erro}")
    desconhecidas = sorted(set(dados) - set(CAMPOS) - set(LISTAS))
    if desconhecidas:
        print(f"Aviso: campos que o modelo não usa: {', '.join(desconhecidas)}")

    doc, conferir = preencher(dados, op.modelo)
    destino = op.saida or SAIDA / f"oficio-{slug(dados.get('assunto'))}.docx"
    n = 2
    while destino.exists() and not op.saida:
        destino = SAIDA / f"oficio-{slug(dados.get('assunto'))}-{n}.docx"
        n += 1
    doc.save(str(destino))
    print(f"Ofício salvo em {destino}")
    print("O bloco de assinatura, o código verificador e o CRC o SEI preenche quando você assina.")
    if conferir:
        print("Para conferir (em amarelo no arquivo):")
        for c in conferir:
            print(f"  - {c}")


if __name__ == "__main__":
    main()
